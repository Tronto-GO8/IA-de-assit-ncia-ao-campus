import os
import re
import fitz
from docx import Document
from pptx import Presentation


PASTA_DOCUMENTOS = "documentos"
PASTA_TEXTOS = "documentos_em_txt"


os.makedirs(PASTA_TEXTOS, exist_ok=True)


# ============================================================
# EXTRAÇÃO
# ============================================================

def extrair_pdf(caminho):
    texto_paginas = []

    with fitz.open(caminho) as pdf:
        for pagina in pdf:
            texto_paginas.append(
                pagina.get_text("text")
            )

    return "\n\n".join(texto_paginas)


def extrair_docx(caminho):
    documento = Document(caminho)

    partes = []

    # Parágrafos
    for paragrafo in documento.paragraphs:
        texto = paragrafo.text.strip()

        if texto:
            partes.append(texto)

    # Tabelas
    for tabela in documento.tables:
        partes.append("\n[TABELA]")

        for linha in tabela.rows:
            celulas = []

            for celula in linha.cells:
                texto = celula.text.strip()
                celulas.append(texto)

            partes.append(" | ".join(celulas))

        partes.append("[/TABELA]")

    return "\n".join(partes)


def extrair_pptx(caminho):
    apresentacao = Presentation(caminho)

    slides = []

    for numero_slide, slide in enumerate(
        apresentacao.slides,
        start=1
    ):
        partes_slide = []

        partes_slide.append(
            f"[SLIDE {numero_slide}]"
        )

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                texto = shape.text.strip()

                if texto:
                    partes_slide.append(texto)

        slides.append(
            "\n".join(partes_slide)
        )

    return "\n\n".join(slides)


def extrair_txt(caminho):
    with open(
        caminho,
        "r",
        encoding="utf-8"
    ) as arquivo:
        return arquivo.read()


# ============================================================
# LIMPEZA
# ============================================================

def limpar_texto(texto):

    # Normaliza quebras de linha
    texto = texto.replace("\r\n", "\n")
    texto = texto.replace("\r", "\n")

    linhas = texto.split("\n")

    linhas_limpas = []

    for linha in linhas:

        linha = linha.strip()

        if not linha:
            linhas_limpas.append("")
            continue

        # Remove linhas que são apenas números
        if linha.isdigit():
            continue

        # Alguns elementos comuns de PDF
        if linha.lower() in [
            "fls",
            "rubrica",
            "sumário",
            "pagina",
            "página"
        ]:
            continue

        linhas_limpas.append(linha)

    texto = "\n".join(linhas_limpas)

    # Espaços duplicados
    texto = re.sub(
        r"[ \t]+",
        " ",
        texto
    )

    # No máximo duas quebras de linha
    texto = re.sub(
        r"\n{3,}",
        "\n\n",
        texto
    )

    return texto.strip()


# ============================================================
# EXTRAÇÃO PRINCIPAL
# ============================================================

def extrair_documento(caminho):

    extensao = os.path.splitext(
        caminho
    )[1].lower()

    if extensao == ".pdf":
        return extrair_pdf(caminho)

    elif extensao == ".docx":
        return extrair_docx(caminho)

    elif extensao == ".pptx":
        return extrair_pptx(caminho)

    elif extensao == ".txt":
        return extrair_txt(caminho)

    return None


# ============================================================
# PROCESSAMENTO DA PASTA
# ============================================================

def processar_documentos():

    arquivos = os.listdir(
        PASTA_DOCUMENTOS
    )

    for arquivo in arquivos:

        extensao = os.path.splitext(
            arquivo
        )[1].lower()

        if extensao not in [
            ".pdf",
            ".docx",
            ".pptx",
            ".txt"
        ]:
            continue

        caminho = os.path.join(
            PASTA_DOCUMENTOS,
            arquivo
        )

        print(
            f"\nProcessando: {arquivo}"
        )

        try:

            texto = extrair_documento(
                caminho
            )

            if not texto:
                print(
                    "  ⚠ Nenhum texto encontrado."
                )
                continue

            texto = limpar_texto(
                texto
            )

            nome_saida = (
                os.path.splitext(arquivo)[0]
                + ".txt"
            )

            caminho_saida = os.path.join(
                PASTA_TEXTOS,
                nome_saida
            )

            with open(
                caminho_saida,
                "w",
                encoding="utf-8"
            ) as arquivo_saida:

                arquivo_saida.write(texto)

            print(
                f"  ✓ Salvo em: {caminho_saida}"
            )

        except Exception as erro:

            print(
                f"  ✗ Erro: {erro}"
            )

    print("\nProcessamento concluído!")


if __name__ == "__main__":
    processar_documentos()